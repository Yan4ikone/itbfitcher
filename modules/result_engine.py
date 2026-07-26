from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from dictionaries.all_dictionaries import MATERIAL_COLORS
from dictionaries.dropdown_lists import DROPDOWN_LISTS


def apply_result(ws, row, code_col, result):
    set_code(ws, row, code_col, result)
    set_dropdown(ws, row, code_col, result)
    set_comment(ws, row, code_col, result)
    set_color(ws, row, code_col, result)

def set_code(ws, row, code_col, result):
    ws.cell(row=row, column=code_col).value = result.code

def set_comment(ws, row, code_col, result):

    if not result.review:
        return

    text = "\n".join(
        f"{code} — {name}"
        for code, name in result.alternatives.items()
    )

    ws.cell(
        row=row,
        column=code_col
    ).comment = Comment(text, "Classifier")

def set_dropdown(ws, row, code_col, result):

    data = DROPDOWN_LISTS.get(result.dropdown)

    if not data:
        return

    codes = [
        item["code"]
        for item in data["variants"]
    ]

    prompt = "\n".join(
        f'{item["code"]} — {item["name"]}'
        for item in data["variants"]
    )

    formula = '"' + ",".join(codes) + '"'

    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=True
    )

    dv.promptTitle = data["title"]
    dv.prompt = prompt
    dv.showInputMessage = True

    cell = f"{get_column_letter(code_col)}{row}"

    dv.add(cell)

    ws.add_data_validation(dv)

def set_color(ws, row, code_col, material_group):

    color = MATERIAL_COLORS.get(material_group)

    if not color:
        return

    ws.cell(row=row, column=code_col).fill = PatternFill("solid", fgColor=color)

