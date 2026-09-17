"""
Разбивка входного файла на куски для нескольких серверов и сборка
результатов обратно в один файл.

Почему можно резать построчно: OzonAutoProcessor обрабатывает каждую
строку независимо (модель — фиксированная ссылка на товар в колонке
"Ссылка" -> DecisionEngine.decide()), кросс-строчной логики между
строками одного файла нет (см. processors/ozon_auto_processor.py::run).
Единственное, что теряется при разбивке — внутрифайловый кэш
"этот же товар уже встречался чуть выше в этом же файле": он просто
не сработает между кусками на разных серверах. На корректность
результата это не влияет, только на скорость (лишний повторный
парсинг одинаковых ссылок в разных кусках).

Формат кусков намеренно НЕ компактный: каждый кусок — это полная
копия исходного файла (те же строки, тот же шаблон, те же макросы),
где у строк ВНЕ диапазона этого куска очищена колонка "Ссылка".
OzonAutoProcessor.get_url_from_row() тогда просто не найдёт там
ссылку и пропустит эти строки — они останутся как были. Это
устраняет любой риск сдвига номеров строк при сборке: строка N в
куске — это всегда строка N в исходном файле.

ВАЖНО (если после split открываешь part-файл в Excel и он "выглядит
как весь исходный файл, только частично пустой"): это ожидаемое
поведение, а не баг. part2.xlsm при --parts 2 на файле из 1000 строк
— это ВСЕ 1000 строк, у которых заполнена колонка "Ссылка" только в
своём диапазоне (например 502-1000), а строки 2-501 в этом файле
выглядят пустыми (ссылка очищена) — это нормально, их обработает
part1 на другом сервере. Нумерация строк Excel (слева) при этом
всегда идёт 1, 2, 3... в любом файле — это не "сброс счётчика", это
просто номера строк листа, они никогда не менялись.

Сборка делает обратное: читает из каждого "_RESULT" куска только
те строки, что реально относились к этому куску, и переносит из них
столбцы B/C/K/L/M/N в мастер-файл — той же строкой. Мастер-файл
никогда не разбирается на части, поэтому макросы/выпадающие списки/
условное форматирование в нём не трогаются вообще.

Есть и третья подкоманда, merge-cards — для card_repository (storage/
runtime_cards.json). "Обучение" (learning/runtime.py::mark_learning_
processed) ищет полную карточку товара именно там по URL, чтобы
перенести её в постоянную storage/knowledge_base.json. Если файл
обрабатывался на сервере A, а "Обучение" запускается на локальной
машине -- локальный storage/runtime_cards.json понятия не имеет об
этой карточке, и "Обучение" её молча пропустит. Поэтому перед
"Обучением" на локальной машине нужно собрать storage/runtime_cards.
json со всех серверов, что участвовали в обработке, и слить их в
один (union по URL -- конфликтов не бывает, если резать split-ом
из этого же файла).

О СКОРОСТИ (почему split/merge могут "долго висеть" без вывода):
openpyxl при обычной загрузке (load_workbook без read_only) читает
в память ВСЕ строки листа вплоть до ws.max_row, а не только реально
заполненные данными. Если в шаблоне форматирование (заливка/шрифт/
границы) применено на тысячи строк вперёд "про запас" - ws.max_row
может быть в 50-100 раз больше реального числа товаров, и именно
это, а не количество товаров, определяет время load()/save(). Ниже
это теперь видно в консоли явно (строка "max_row=... реальных данных
до строки..."), и на файлах-результатах (только для merge, split
всё ещё должен уметь ПИСАТЬ, поэтому грузится в обычном режиме)
используется read_only-режим с последовательным чтением построчно
(iter_rows), а не поштучный доступ к ячейкам вроде ws["E123"] -
поштучный доступ в read_only-режиме на больших листах на порядки
МЕДЛЕННЕЕ обычной загрузки (каждое такое обращение пересканирует
лист заново), поэтому здесь важно, что чтение именно
последовательное.

Использование:

    python tools/fleet_split_merge.py split input.xlsm --parts 4
        -> input_part1.xlsm ... input_part4.xlsm

    python tools/fleet_split_merge.py merge input.xlsm \\
        input_part1_RESULT.xlsm input_part2_RESULT.xlsm ... \\
        --output input_MERGED.xlsm

    python tools/fleet_split_merge.py merge-cards \\
        storage/runtime_cards.json \\
        server1_runtime_cards.json server2_runtime_cards.json ... \\
        --output storage/runtime_cards.json
"""

import argparse
import json
import sys
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

# Колонки жёстко совпадают с processors/ozon_auto_processor.py —
# если там когда-нибудь поменяется схема, поменять и здесь.
LINK_COL = "E"          # OzonAutoProcessor.get_url_from_row проверяет D и E,
                         # но реальный текстовый URL лежит в E (D = формула
                         # =HYPERLINK(E2), с ней get_url_from_row не сработает)
DESCRIPTION_COL = "B"
CODE_COL = "C"
MERGE_COLUMNS = ["B", "C", "K", "L", "M", "N"]


def _fmt_elapsed(seconds):
    """Человекочитаемое время для прогресс-логов."""

    if seconds < 60:
        return f"{seconds:.1f} сек"

    minutes = int(seconds // 60)
    secs = seconds - minutes * 60
    return f"{minutes} мин {secs:.0f} сек"


def _report_bloat(ws, last_row, label):
    """См. блок "О СКОРОСТИ" в docstring файла: если реальных данных
    сильно меньше, чем ws.max_row, именно это (пустое форматирование
    "про запас") - главная причина долгой загрузки/сохранения, а не
    число товаров в файле. Печатаем это явно, чтобы было видно, что
    программа не зависла, а действительно читает большой файл."""

    extra = ws.max_row - last_row

    print(
        f"  [{label}] max_row={ws.max_row}, реальных данных до строки "
        f"{last_row}"
        + (
            f" -> {extra} пустых отформатированных строк сверху "
            "(вероятная причина долгой загрузки/сохранения)"
            if extra > 500 else ""
        )
    )


def find_last_data_row(ws):
    """Та же логика, что и в OzonAutoProcessor.find_last_data_row —
    ws.max_row в openpyxl часто завышен из-за форматирования,
    применённого на тысячи строк вперёд в шаблоне.

    ВАЖНО: рассчитана на обычный (не read_only) режим openpyxl, где
    поштучный доступ к ячейкам быстрый. В read_only-режиме такой
    поиск с конца листа может быть очень медленным на раздутых
    файлах — для read_only-сценариев (см. cmd_merge) last_row
    считается один раз с мастер-файла и переиспользуется для всех
    result-файлов, а не пересчитывается для каждого заново."""

    for row in range(ws.max_row, 1, -1):
        if str(ws[f"B{row}"].value or "").strip():
            return row
        if str(ws[f"C{row}"].value or "").strip():
            return row
        value = ws[f"{LINK_COL}{row}"].value
        if isinstance(value, str) and value.startswith("http"):
            return row

    return 1


def row_ranges(first_row, last_row, parts):
    """Делит [first_row, last_row] на `parts` примерно равных
    непрерывных диапазонов (включительно с обеих сторон)."""

    total = last_row - first_row + 1
    if total <= 0:
        return []

    base = total // parts
    remainder = total % parts

    ranges = []
    current = first_row

    for i in range(parts):
        size = base + (1 if i < remainder else 0)
        if size == 0:
            continue
        ranges.append((current, current + size - 1))
        current += size

    return ranges


def cmd_split(args):
    input_path = Path(args.input)

    if not input_path.exists():
        sys.exit(f"Файл не найден: {input_path}")

    t_total_start = time.perf_counter()

    print(f"Загружаю {input_path.name}...")
    t0 = time.perf_counter()
    wb = openpyxl.load_workbook(input_path, keep_vba=True)
    ws = wb.active
    print(f"  загружено за {_fmt_elapsed(time.perf_counter() - t0)}")

    last_row = find_last_data_row(ws)
    print(f"Последняя строка с данными: {last_row}")
    _report_bloat(ws, last_row, label="исходный файл")

    ranges = row_ranges(2, last_row, args.parts)

    if not ranges:
        sys.exit("Нет данных для разбивки (пустой файл?)")

    for i, (start, end) in enumerate(ranges, start=1):

        t_part_start = time.perf_counter()
        print(f"Собираю part{i}/{len(ranges)} (строки {start}-{end})...")

        part_wb = openpyxl.load_workbook(input_path, keep_vba=True)
        part_ws = part_wb.active

        cleared = 0
        for row in range(2, last_row + 1):
            if start <= row <= end:
                continue
            cell = part_ws[f"{LINK_COL}{row}"]
            if cell.value:
                cell.value = None
                cleared += 1

        out_path = input_path.with_name(
            f"{input_path.stem}_part{i}{input_path.suffix}"
        )
        part_wb.save(out_path)

        print(
            f"  part{i}: строки {start}-{end} "
            f"({end - start + 1} шт.), очищено вне диапазона: {cleared} "
            f"-> {out_path.name} "
            f"[{_fmt_elapsed(time.perf_counter() - t_part_start)}]"
        )

    print(
        "\nГотово за {total}. Раздайте part1..part{n} по серверам, "
        "соберите ...RESULT.xlsm обратно командой merge.".format(
            n=len(ranges),
            total=_fmt_elapsed(time.perf_counter() - t_total_start),
        )
    )


def cmd_merge(args):
    master_path = Path(args.master)

    if not master_path.exists():
        sys.exit(f"Мастер-файл не найден: {master_path}")

    t_total_start = time.perf_counter()

    print(f"Загружаю мастер-файл {master_path.name}...")
    t0 = time.perf_counter()
    master_wb = openpyxl.load_workbook(master_path, keep_vba=True)
    master_ws = master_wb.active
    print(f"  загружено за {_fmt_elapsed(time.perf_counter() - t0)}")

    # Считаем last_row ОДИН раз с мастер-файла и переиспользуем для
    # всех result-файлов ниже, вместо пересчёта на каждом из них.
    # Это безопасно по самому дизайну split: каждый result-файл -
    # это полная копия того же исходного файла (см. docstring), так
    # что строка N значит одно и то же везде. Заодно это позволяет
    # читать result-файлы в read_only-режиме без медленного поиска
    # последней строки с конца листа (см. предупреждение в
    # find_last_data_row).
    master_last_row = find_last_data_row(master_ws)
    print(f"Последняя строка с данными в мастер-файле: {master_last_row}")
    _report_bloat(master_ws, master_last_row, label="мастер-файл")

    col_positions = {
        col: column_index_from_string(col) - 1
        for col in ([LINK_COL] + MERGE_COLUMNS)
    }
    max_col = max(col_positions.values()) + 1

    total_transferred = 0

    for result_file in args.results:
        result_path = Path(result_file)

        if not result_path.exists():
            print(f"  ПРОПУСК (не найден): {result_path}")
            continue

        t_file_start = time.perf_counter()

        # read_only=True: этот файл нам нужен ТОЛЬКО на чтение,
        # обратно в него ничего не сохраняем - обычный режим тут
        # означал бы разбор ВСЕГО раздутого листа впустую (см.
        # docstring, блок "О СКОРОСТИ"). Важно: читаем строго
        # последовательно через iter_rows(), а НЕ поштучным доступом
        # вида result_ws["E123"] - в read_only-режиме поштучный
        # доступ на больших листах в десятки раз медленнее обычной
        # загрузки, а не быстрее.
        result_wb = openpyxl.load_workbook(
            result_path, read_only=True, data_only=True
        )
        result_ws = result_wb.active

        transferred = 0

        for row_idx, row_cells in enumerate(
            result_ws.iter_rows(
                min_row=2, max_row=master_last_row, max_col=max_col
            ),
            start=2,
        ):
            link_value = row_cells[col_positions[LINK_COL]].value

            # Строка не относилась к этому куску (ссылку мы сами
            # очистили на split) -- пропускаем, чтобы не затереть
            # мастер-файл пустотой.
            if not link_value:
                continue

            for col in MERGE_COLUMNS:
                value = row_cells[col_positions[col]].value
                if value not in (None, ""):
                    master_ws[f"{col}{row_idx}"] = value

            transferred += 1

        result_wb.close()

        print(
            f"  {result_path.name}: перенесено строк: {transferred} "
            f"[{_fmt_elapsed(time.perf_counter() - t_file_start)}]"
        )
        total_transferred += transferred

    output_path = Path(args.output) if args.output else master_path.with_name(
        f"{master_path.stem}_MERGED{master_path.suffix}"
    )

    print(f"Сохраняю {output_path.name}...")
    t0 = time.perf_counter()
    master_wb.save(output_path)
    print(f"  сохранено за {_fmt_elapsed(time.perf_counter() - t0)}")

    print(f"\nГотово за {_fmt_elapsed(time.perf_counter() - t_total_start)}.")
    print(f"Всего перенесено строк: {total_transferred}")
    print(f"Сохранено: {output_path}")


def cmd_merge_cards(args):
    t_total_start = time.perf_counter()
    merged = {}
    total_before = 0

    for source_file in [args.local] + args.servers:
        source_path = Path(source_file)

        if not source_path.exists():
            print(f"  ПРОПУСК (не найден): {source_path}")
            continue

        with open(source_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        total_before += len(data)
        overwritten = sum(1 for url in data if url in merged)
        merged.update(data)

        print(
            f"  {source_path.name}: {len(data)} карточек "
            f"({overwritten} перезаписали более раннюю версию по тому же URL)"
        )

    output_path = Path(args.output) if args.output else args.local
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(merged, f, ensure_ascii=False, indent=4)

    print(
        f"\nГотово за {_fmt_elapsed(time.perf_counter() - t_total_start)}. "
        f"Уникальных карточек после слияния: {len(merged)} "
        f"(прочитано всего записей: {total_before})"
    )
    print(f"Сохранено: {output_path}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    split_parser = subparsers.add_parser(
        "split", help="Разбить входной файл на N кусков для серверов"
    )
    split_parser.add_argument("input", help="Путь к входному .xlsm")
    split_parser.add_argument(
        "--parts", type=int, required=True, help="На сколько серверов делим"
    )
    split_parser.set_defaults(func=cmd_split)

    merge_parser = subparsers.add_parser(
        "merge", help="Собрать результаты с серверов обратно в один файл"
    )
    merge_parser.add_argument(
        "master",
        help="Исходный (нераздробленный) файл -- в него переносятся результаты",
    )
    merge_parser.add_argument(
        "results", nargs="+", help="Файлы partN_RESULT.xlsm со всех серверов"
    )
    merge_parser.add_argument(
        "--output", help="Куда сохранить итог (по умолчанию: <master>_MERGED.xlsm)"
    )
    merge_parser.set_defaults(func=cmd_merge)

    cards_parser = subparsers.add_parser(
        "merge-cards",
        help="Слить storage/runtime_cards.json со всех серверов перед Обучением",
    )
    cards_parser.add_argument(
        "local", help="Локальный runtime_cards.json (или путь для нового файла)"
    )
    cards_parser.add_argument(
        "servers", nargs="+", help="runtime_cards.json, собранные с серверов"
    )
    cards_parser.add_argument(
        "--output", help="Куда сохранить итог (по умолчанию: перезаписать local)"
    )
    cards_parser.set_defaults(func=cmd_merge_cards)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()